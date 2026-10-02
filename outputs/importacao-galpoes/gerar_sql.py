import datetime
import hashlib
import json
import pathlib
import re

import openpyxl

source = pathlib.Path(r'C:\Users\vchaves\Downloads\galpoes.xlsx')
destination = pathlib.Path(__file__).with_name('galpoes_insert_estrutura_atual.sql')


def normalize(value):
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip() or None


def quote(value):
    return 'NULL' if value is None else "'" + value.replace("'", "''") + "'"


values = []
hashes = set()
for sheet in openpyxl.load_workbook(source, data_only=True):
    assert sheet.cell(1, 2).value == 'Granja'
    assert sheet.cell(1, 5).value == 'Galpão'
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        full = [normalize(value) for value in row]
        if not any(full):
            continue
        assert re.match(r'^\d+\s*-', full[1] or ''), (sheet.title, index)
        assert full[4], (sheet.title, index)
        payload = {'version': 'galpoes-xlsx-v1', 'aba': sheet.title, 'colunas': full}
        digest = hashlib.sha256(json.dumps(payload, ensure_ascii=False, separators=(',', ':')).encode('utf-8')).hexdigest()
        assert digest not in hashes
        hashes.add(digest)
        fields = [sheet.title, source.name, digest, full[1], full[4], full[7], full[8], None, full[6]]
        values.append('    (' + ', '.join(quote(value) for value in fields) + ')')

assert len(values) == 1869
sql = """-- IMPORTACAO PARCIAL: apenas as colunas existentes informadas.
-- Conferir schema zootecnico, empresa 1 e unidade 52 antes de executar.
-- Modelo, Tipo Granja e datas de atividade NAO sao gravados.
-- ramo_atividade ausente na planilha: NULL.
-- hash_linha: SHA-256 da linha completa normalizada, formato galpoes-xlsx-v1.
-- Evita repeticoes sequenciais deste arquivo por hash e escopo.
-- Nao garante deduplicacao contra hashes produzidos por outros importadores.
BEGIN;
SET LOCAL standard_conforming_strings = on;
WITH parametros(empresa_codigo, unidade_codigo) AS (
    VALUES ('1'::text, '52'::text)
), dados(aba, origem_arquivo, hash_linha, granja, galpao, tecnico,
          supervisor, ramo_atividade, municipio) AS (
    VALUES
""" + ',\n'.join(values) + """
), inseridos AS (
    INSERT INTO zootecnico.galpoes (
        empresa_codigo, unidade_codigo, aba, origem_arquivo, hash_linha,
        granja, galpao, tecnico, supervisor, ramo_atividade, municipio
    )
    SELECT p.empresa_codigo, p.unidade_codigo, d.aba, d.origem_arquivo,
           d.hash_linha, d.granja, d.galpao, d.tecnico, d.supervisor,
           d.ramo_atividade, d.municipio
    FROM dados d CROSS JOIN parametros p
    WHERE NOT EXISTS (
        SELECT 1 FROM zootecnico.galpoes g
        WHERE g.empresa_codigo = p.empresa_codigo
          AND g.unidade_codigo = p.unidade_codigo
          AND g.hash_linha = d.hash_linha
    )
    RETURNING id
)
SELECT COUNT(*) AS registros_inseridos FROM inseridos;
COMMIT;
"""
# Importacao parcial aposentada; o gerador completo esta no arquivo separado.
if __name__ == '__main__':
    import runpy
    runpy.run_path(str(pathlib.Path(__file__).with_name('gerar_sql_completo.py')), run_name='__main__')
