import collections
import datetime
import hashlib
import json
import pathlib
import re
import unicodedata
import openpyxl
from openpyxl.utils import get_column_letter

source = pathlib.Path(r'C:\Users\vchaves\Downloads\galpoes.xlsx')
directory = pathlib.Path(__file__).parent
book = openpyxl.load_workbook(source, data_only=True)
sheet = book.active

def normalize(v):
    if v is None: return None
    if isinstance(v, (datetime.datetime, datetime.date)): return v.isoformat()
    if isinstance(v, (int, float)) and float(v).is_integer(): return str(int(v))
    return str(v).strip() or None

def quote(v):
    return 'NULL' if v is None else "'" + v.replace("'", "''") + "'"

def identifier(v):
    raw = re.sub(r'_x[0-9a-fA-F]{4}_', ' ', str(v))
    raw = ''.join(c for c in unicodedata.normalize('NFKD', raw) if not unicodedata.combining(c))
    raw = re.sub(r'[^a-z0-9]+', '_', raw.lower()).strip('_')
    return 'campo_' + raw if raw[0].isdigit() else raw

columns, mapping, used = [], [], collections.Counter()
for cell in sheet[1]:
    observed = [sheet.cell(r, cell.column).value for r in range(2, sheet.max_row+1) if sheet.cell(r, cell.column).value is not None]
    if cell.value is None and not observed: continue
    base = identifier(cell.value) if cell.value is not None else 'coluna_' + get_column_letter(cell.column).lower()
    used[base] += 1
    name = base if used[base] == 1 else f'{base}_{used[base]}'
    assert len(name) <= 63
    date = any(isinstance(v, (datetime.date, datetime.datetime)) for v in observed)
    if date: assert all(isinstance(v, (datetime.date, datetime.datetime)) for v in observed)
    kind = 'timestamp without time zone' if date else 'text'
    columns.append((cell.column-1, name, kind))
    mapping.append({'excel':get_column_letter(cell.column), 'cabecalho':cell.value, 'coluna':name, 'tipo':kind})
names = [n for _,n,_ in columns]
assert len(names) == len(set(names)) == 135
ddl = '-- Preserva colunas e dados existentes. Executar antes da importacao.\nBEGIN;\nALTER TABLE zootecnico.galpoes\n'
ddl += ',\n'.join(f'    ADD COLUMN IF NOT EXISTS {n} {t}' for _,n,t in columns if n not in {'granja','galpao','supervisor'})
(directory/'01_ampliar_colunas_galpoes.sql').write_text(ddl+';\nCOMMIT;\n',encoding='utf-8')
metadata = ['empresa_codigo','unidade_codigo','aba','origem_arquivo','hash_linha']
stage = metadata + names
definitions = [n+' text' for n in metadata] + [n+' '+t for _,n,t in columns]
values, hashes = [], set()
for s in book:
    assert [c.value for c in s[1]] == [c.value for c in sheet[1]]
    for row in s.iter_rows(min_row=2,values_only=True):
        full = [normalize(v) for v in row]
        if not any(full): continue
        assert re.match(r'^\d+\s*-',full[1] or '') and full[4]
        payload = {'version':'galpoes-xlsx-v1','aba':s.title,'colunas':full}
        h = hashlib.sha256(json.dumps(payload,ensure_ascii=False,separators=(',',':')).encode('utf-8')).hexdigest()
        assert h not in hashes
        hashes.add(h)
        fields = ['1','52',s.title,source.name,h] + [full[i] for i,_,_ in columns]
        values.append('    ('+', '.join(quote(v) for v in fields)+')')
assert len(values) == 1869
target = ['aba','origem_arquivo'] + names + ['tecnico','municipio']
expr = ['s.'+n for n in target[:-2]] + ['s.tec','s.cidade']
join = 'g.empresa_codigo = s.empresa_codigo AND g.unidade_codigo = s.unidade_codigo AND g.hash_linha = s.hash_linha'
sql = '''-- Executar 01_ampliar_colunas_galpoes.sql primeiro.
-- Escopo empresa 1/unidade 52: conferir e editar no UPDATE de escopo abaixo.
-- Hash compativel com o SQL parcial anterior desta conversa.
-- Datas tipadas; demais campos text, sem inventar tipos para dados mistos.
BEGIN;
SET LOCAL standard_conforming_strings = on;
CREATE TEMP TABLE carga_galpoes_completa (
''' + ',\n'.join('    '+d for d in definitions) + '\n) ON COMMIT DROP;\n'
sql += 'INSERT INTO carga_galpoes_completa ('+', '.join(stage)+') VALUES\n'+',\n'.join(values)+';\n'
sql += '''
-- Edite SOMENTE estes dois valores se o escopo for diferente.
UPDATE carga_galpoes_completa SET empresa_codigo = '1', unidade_codigo = '52';
DO $$
BEGIN
    IF EXISTS (
        SELECT 1 FROM zootecnico.galpoes g
        JOIN carga_galpoes_completa s ON '''+join+'''
        GROUP BY g.empresa_codigo, g.unidade_codigo, g.hash_linha
        HAVING COUNT(*) > 1
    ) THEN
        RAISE EXCEPTION 'Hashes duplicados no banco: revisar antes de importar.';
    END IF;
END $$;
WITH atualizados AS (
    UPDATE zootecnico.galpoes g SET
'''+',\n'.join('        '+n+' = '+e for n,e in zip(target,expr))
sql += '\n    FROM carga_galpoes_completa s\n    WHERE '+join
sql += '\n    AND ROW('+', '.join('g.'+n for n in target)+') IS DISTINCT FROM ROW('+', '.join(expr)+')\n    RETURNING g.id\n)\nSELECT COUNT(*) AS registros_atualizados FROM atualizados;\n'
sql += '\nWITH inseridos AS (\n    INSERT INTO zootecnico.galpoes ('+', '.join(['empresa_codigo','unidade_codigo','hash_linha']+target)+')\n    SELECT '+', '.join(['s.empresa_codigo','s.unidade_codigo','s.hash_linha']+expr)
sql += '\n    FROM carga_galpoes_completa s\n    WHERE NOT EXISTS (SELECT 1 FROM zootecnico.galpoes g WHERE '+join+')\n    RETURNING id\n)\nSELECT COUNT(*) AS registros_inseridos FROM inseridos;\nCOMMIT;\n'
(directory/'02_importar_galpoes_completo.sql').write_text(sql,encoding='utf-8')
(directory/'mapeamento_colunas_galpoes.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
(directory/'galpoes_insert_estrutura_atual.sql').write_text('-- SUBSTITUIDO: executar 01_ampliar_colunas_galpoes.sql e depois 02_importar_galpoes_completo.sql.\n',encoding='utf-8')
print(json.dumps({'linhas':len(values),'colunas':len(columns),'datas':[n for _,n,t in columns if t!='text'],'repetidos':{k:v for k,v in used.items() if v>1}},ensure_ascii=True))
